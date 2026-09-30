"""月間週次予定表(実績表)の合計チェック。

各週の明細行を足し上げて【第N週 合計】と一致するか、
週合計の和が【月合計】と一致するかを「工数」「うち残業」の両方で確認する。

使い方: python3 月ノルマ/check_totals.py 月ノルマ/月間週次予定表_2026年10月版.xlsx
"""
import re
import sys

import openpyxl

COL_NAME, COL_HOURS, COL_OT = 2, 5, 6  # B, E, F


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def main(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    ok = True
    detail = [0.0, 0.0]
    weeks = [0.0, 0.0]
    for r in range(1, ws.max_row + 1):
        name = str(ws.cell(r, COL_NAME).value or "")
        h, ot = num(ws.cell(r, COL_HOURS).value), num(ws.cell(r, COL_OT).value)
        if re.search(r"【第\d週\s*合計】", name):
            match = (h, ot) == tuple(detail)
            ok &= match
            print(f"{'OK' if match else 'NG'} {name}: 明細 {detail[0]}h/残業{detail[1]}h → 記載 {h}h/残業{ot}h")
            weeks = [weeks[0] + h, weeks[1] + ot]
            detail = [0.0, 0.0]
        elif re.search(r"【\d+月\s*月合計】", name):
            match = (h, ot) == tuple(weeks)
            ok &= match
            print(f"{'OK' if match else 'NG'} {name}: 週合計の和 {weeks[0]}h/残業{weeks[1]}h → 記載 {h}h/残業{ot}h")
        elif name and r > 7:
            detail = [detail[0] + h, detail[1] + ot]
    print("最終チェック:", "一致" if ok else "不一致あり")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
