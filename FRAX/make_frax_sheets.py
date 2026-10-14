"""FRAX対象者の抽出とFRAXシート作成。

患者一覧(xlsx の「患者一覧」シート、または健診システムから出力した CSV)を読み、
「設定」シートの条件で FRAX 評価の対象になる患者だけを抽出して、
1人1シートの FRAX シートと判定一覧を持つ Excel を作る。

使い方:
  python3 FRAX/make_frax_sheets.py --template FRAX/FRAX入力_テンプレート.xlsx
      … 入力用テンプレート(患者一覧・設定シート、架空のサンプル患者入り)を作る
  python3 FRAX/make_frax_sheets.py <入力.xlsx|入力.csv> [出力.xlsx]
      … 対象者だけの FRAX シートを作る(出力省略時は「FRAXシート_<入力名>.xlsx」)

判定の既定値(骨粗鬆症の予防と治療ガイドライン2015年版の薬物治療開始基準に準拠):
  - 年齢 40〜90歳(FRAX の適用範囲)
  - 椎体・大腿骨近位部の脆弱性骨折あり → 骨粗鬆症と診断されるため FRAX 不要(対象外)
  - その他の脆弱性骨折あり かつ YAM 80%未満 → 骨粗鬆症(対象外)
  - YAM 70%以下 → 骨粗鬆症(対象外)/ YAM 80%以上 → 正常域(対象外)
  - YAM 70%超〜80%未満(骨量減少)→ FRAX 対象
CSV で条件を変えたいときは、テンプレートの「設定」シートを編集して xlsx で入力する。

FRAX の確率そのものは公式サイト(https://frax.shef.ac.uk/FRAX/tool.aspx?lang=jp)で計算し、
各シートの黄色セルに入力すると判定(15%以上で薬物治療開始を検討)が自動表示される。
"""
import csv
import datetime as dt
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# 入力列(見出し → 内部キー)。見出しは別名も受け付ける。
COLUMNS = [
    ("受診日", "date", ["健診日", "検査日"]),
    ("患者ID", "pid", ["ID", "カルテ番号", "受診者番号"]),
    ("氏名", "name", ["患者名", "受診者名"]),
    ("性別", "sex", []),
    ("生年月日", "birth", []),
    ("年齢", "age", []),
    ("身長(cm)", "height", ["身長"]),
    ("体重(kg)", "weight", ["体重"]),
    ("骨密度YAM(%)", "yam", ["YAM", "YAM(%)", "YAM値"]),
    ("大腿骨頸部骨密度(g/cm2)", "bmd", ["大腿骨頸部BMD", "BMD", "骨密度(g/cm2)"]),
    ("骨密度測定部位", "site", ["測定部位"]),
    ("椎体・大腿骨近位部の骨折歴", "major_fx", ["椎体骨折", "大腿骨近位部骨折"]),
    ("その他の脆弱性骨折歴", "other_fx", ["骨折歴", "脆弱性骨折"]),
    ("両親の大腿骨近位部骨折歴", "parent_fx", ["両親の大腿骨骨折歴"]),
    ("現在の喫煙", "smoke", ["喫煙"]),
    ("糖質コルチコイド", "steroid", ["ステロイド", "ステロイド使用"]),
    ("関節リウマチ", "ra", ["リウマチ"]),
    ("続発性骨粗鬆症", "secondary", []),
    ("アルコール1日3単位以上", "alcohol", ["アルコール", "飲酒"]),
]
YES_NO_KEYS = ["major_fx", "other_fx", "parent_fx", "smoke", "steroid", "ra", "secondary", "alcohol"]

# 設定(項目名 → 既定値, 説明)
SETTINGS = [
    ("年齢下限", 40, "FRAX の適用範囲は40〜90歳"),
    ("年齢上限", 90, ""),
    ("対象性別", "全員", "全員 / 女性 / 男性"),
    ("YAM下限(この値を超える)", 70, "70%以下は骨粗鬆症と診断されるため FRAX 不要"),
    ("YAM上限(この値未満)", 80, "80%以上は正常域"),
    ("椎体・大腿骨近位部骨折ありは除外", "はい", "骨折があれば骨粗鬆症と診断されるため"),
    ("その他の脆弱性骨折+YAM上限未満は除外", "はい", "同上(骨粗鬆症と診断)"),
    ("YAM未測定の扱い", "要確認", "要確認(シートを作らない)/ 対象(骨密度なしで FRAX を計算)"),
    ("薬物治療検討の基準(主要骨折10年確率 %)", 15, "ガイドライン2015:75歳未満で15%以上"),
    ("基準を適用する年齢(未満)", 75, "75歳以上は15%基準を一律に当てはめない"),
]

NAVY, BLUE, YELLOW, GRAY = "17324D", "C7D9E6", "FFF2CC", "F2F2F2"
TARGET_FILL, EXCLUDE_FILL, CHECK_FILL = "EAF3EA", "F2F2F2", "FCE4D6"
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def fill(color):
    return PatternFill("solid", fgColor=color)


def norm_header(s):
    return re.sub(r"[\s()()]", "", str(s or ""))


def header_map(headers):
    lookup = {}
    for label, key, aliases in COLUMNS:
        for name in [label, *aliases]:
            lookup.setdefault(norm_header(name), key)
    return {i: lookup[norm_header(h)] for i, h in enumerate(headers) if norm_header(h) in lookup}


def to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"-?\d+(\.\d+)?", str(v).translate(str.maketrans("0123456789.", "0123456789.")))
    return float(m.group()) if m else None


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if not v:
        return None
    s = str(v).strip().replace("/", "-").replace(".", "-")
    m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", s) or re.match(r"(\d{4})年(\d{1,2})月(\d{1,2})日", str(v))
    return dt.date(*map(int, m.groups())) if m else None


def yes(v):
    return str(v or "").strip().lower() in {"はい", "あり", "有", "yes", "y", "1", "○", "〇", "true", "+"}


def sex_of(v):
    s = str(v or "").strip()
    if s in {"女", "女性", "F", "f", "2"}:
        return "女性"
    if s in {"男", "男性", "M", "m", "1"}:
        return "男性"
    return s


def age_on(birth, on):
    return on.year - birth.year - ((on.month, on.day) < (birth.month, birth.day))


def read_rows(path):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        for enc in ("utf-8-sig", "cp932"):
            try:
                with open(path, encoding=enc, newline="") as f:
                    rows = list(csv.reader(f))
                break
            except UnicodeDecodeError:
                continue
        return rows, {}
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["患者一覧"] if "患者一覧" in wb.sheetnames else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    settings = {}
    if "設定" in wb.sheetnames:
        for r in wb["設定"].iter_rows(min_row=2, values_only=True):
            if r and r[0]:
                settings[str(r[0]).strip()] = r[1]
    return rows, settings


def load_patients(rows):
    # 見出し行は、既知の列名が3つ以上ある最初の行
    for hi, r in enumerate(rows):
        cols = header_map(r)
        if len(cols) >= 3:
            break
    else:
        sys.exit("見出し行が見つかりません(患者ID・氏名・性別などの列名が必要です)")
    patients = []
    for r in rows[hi + 1:]:
        p = {key: (r[i] if i < len(r) else None) for i, key in cols.items()}
        if not any(v not in (None, "") for v in p.values()):
            continue
        patients.append(p)
    return patients


def get_settings(user):
    s = {name: default for name, default, _ in SETTINGS}
    for name in s:
        if user.get(name) not in (None, ""):
            s[name] = user[name]
    return s


def judge(p, s):
    """(区分, 理由) を返す。区分は 対象 / 対象外 / 要確認。"""
    on = to_date(p.get("date")) or dt.date.today()
    age = to_num(p.get("age"))
    if age is None and to_date(p.get("birth")):
        age = age_on(to_date(p.get("birth")), on)
    p["_age"] = int(age) if age is not None else None
    p["_sex"] = sex_of(p.get("sex"))
    yam = to_num(p.get("yam"))

    if age is None:
        return "要確認", "年齢・生年月日が未入力"
    lo, hi = to_num(s["年齢下限"]), to_num(s["年齢上限"])
    if not lo <= age <= hi:
        return "対象外", f"{int(age)}歳(FRAX適用は{lo:g}〜{hi:g}歳)"
    want = str(s["対象性別"]).strip()
    if want in {"女性", "男性"} and p["_sex"] != want:
        return "対象外", f"対象性別外({p['_sex'] or '未入力'})"
    if p["_sex"] not in {"女性", "男性"}:
        return "要確認", "性別が未入力・不明"
    if yes(s["椎体・大腿骨近位部骨折ありは除外"]) and yes(p.get("major_fx")):
        return "対象外", "椎体・大腿骨近位部骨折あり→骨粗鬆症(治療検討へ)"
    ylo, yhi = to_num(s["YAM下限(この値を超える)"]), to_num(s["YAM上限(この値未満)"])
    if yam is None:
        if str(s["YAM未測定の扱い"]).strip() == "対象":
            return "対象", "YAM未測定(骨密度なしで計算)"
        return "要確認", "骨密度(YAM)未測定"
    if yes(s["その他の脆弱性骨折+YAM上限未満は除外"]) and yes(p.get("other_fx")) and yam < yhi:
        return "対象外", f"脆弱性骨折あり+YAM{yam:g}%→骨粗鬆症(治療検討へ)"
    if yam <= ylo:
        return "対象外", f"YAM{yam:g}%→骨粗鬆症域(治療検討へ)"
    if yam >= yhi:
        return "対象外", f"YAM{yam:g}%→正常域"
    return "対象", f"YAM{yam:g}%(骨量減少域)"


def sheet_title(p, used):
    base = re.sub(r"[\\/*?:\[\]]", "", f"FRAX_{p.get('pid') or ''}_{p.get('name') or ''}")[:31].rstrip("_")
    title, n = base, 2
    while title in used:
        title = f"{base[:28]}_{n}"
        n += 1
    used.add(title)
    return title


def write_frax_sheet(ws, p, s, reason):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 44
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.merge_cells("A1:D1")
    ws["A1"] = "FRAX 骨折リスク評価シート"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = fill(NAVY)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    info = [
        ("受診日", to_date(p.get("date"))),
        ("患者ID", p.get("pid")),
        ("氏名", p.get("name")),
        ("抽出理由", reason),
    ]
    r = 3
    for label, val in info:
        ws.cell(r, 2, label).fill = fill(GRAY)
        ws.cell(r, 3, val)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        if isinstance(val, dt.date):
            ws.cell(r, 3).number_format = "yyyy/mm/dd"
        for c in (2, 3, 4):
            ws.cell(r, c).border = BOX
        r += 1

    r += 1
    for c, h in enumerate(["No", "FRAX 入力項目", "値", "備考"], 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(bold=True)
        cell.fill = fill(BLUE)
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center")
    r += 1

    yn = lambda k: "はい" if yes(p.get(k)) else "いいえ"
    fx = "はい" if yes(p.get("major_fx")) or yes(p.get("other_fx")) else "いいえ"
    items = [
        ("年齢", p["_age"], "歳"),
        ("性別", p["_sex"], ""),
        ("体重(kg)", to_num(p.get("weight")), ""),
        ("身長(cm)", to_num(p.get("height")), ""),
        ("BMI", None, "体重・身長から自動計算"),
        ("骨折歴(成人後の脆弱性骨折)", fx, ""),
        ("両親の大腿骨近位部骨折歴", yn("parent_fx"), ""),
        ("現在の喫煙", yn("smoke"), ""),
        ("糖質コルチコイド", yn("steroid"), "3か月以上・PSL換算5mg/日以上"),
        ("関節リウマチ", yn("ra"), ""),
        ("続発性骨粗鬆症", yn("secondary"), "1型糖尿病・甲状腺機能亢進症・早期閉経など"),
        ("アルコール(1日3単位以上)", yn("alcohol"), "1単位=アルコール8〜10g"),
        ("大腿骨頸部骨密度(g/cm2)", to_num(p.get("bmd")), f"YAM {to_num(p.get('yam')) if to_num(p.get('yam')) is not None else '未測定'}%"
         + (f"/測定部位:{p.get('site')}" if p.get("site") else "")),
    ]
    first = r
    for i, (label, val, note) in enumerate(items, 1):
        ws.cell(r, 1, i).alignment = Alignment(horizontal="center")
        ws.cell(r, 2, label)
        ws.cell(r, 3, val)
        ws.cell(r, 4, note).font = Font(size=9, color="555555")
        for c in range(1, 5):
            ws.cell(r, c).border = BOX
        if val in (None, "") and label != "BMI":
            ws.cell(r, 3).fill = fill(YELLOW)
        r += 1
    w, h, bmi = f"C{first + 2}", f"C{first + 3}", f"C{first + 4}"
    ws[bmi] = f'=IF(AND(ISNUMBER({w}),ISNUMBER({h}),{h}>0),ROUND({w}/({h}/100)^2,1),"")'
    yn_dv = DataValidation(type="list", formula1='"はい,いいえ"', allow_blank=True)
    ws.add_data_validation(yn_dv)
    yn_dv.add(f"C{first + 5}:C{first + 11}")

    r += 1
    ws.cell(r, 2, "FRAX 計算結果(公式サイトで計算して入力)").font = Font(bold=True)
    r += 1
    res = []
    for label in ("主要骨粗鬆症性骨折 10年確率(%)", "大腿骨近位部骨折 10年確率(%)"):
        ws.cell(r, 2, label)
        ws.cell(r, 3).fill = fill(YELLOW)
        for c in (2, 3):
            ws.cell(r, c).border = BOX
        res.append(f"C{r}")
        r += 1

    th, age_lim = to_num(s["薬物治療検討の基準(主要骨折10年確率 %)"]), to_num(s["基準を適用する年齢(未満)"])
    age_cell = f"C{first}"
    ws.cell(r, 2, "判定").font = Font(bold=True)
    ws.cell(r, 3, (
        f'=IF({res[0]}="","FRAX計算後に結果を入力",'
        f'IF({age_cell}>={age_lim:g},"{age_lim:g}歳以上:{th:g}%基準は適用外(医師の総合判断)",'
        f'IF({res[0]}>={th:g},"{th:g}%以上:薬物治療開始を検討","{th:g}%未満:生活指導・経過観察")))'
    ))
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    ws.cell(r, 3).font = Font(bold=True, color="C00000")
    for c in (2, 3, 4):
        ws.cell(r, c).border = BOX
    r += 2

    for label in ("評価者", "確認医師", "記入日"):
        ws.cell(r, 2, label).fill = fill(GRAY)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        for c in (2, 3, 4):
            ws.cell(r, c).border = BOX
        r += 1
    r += 1
    ws.cell(r, 2, "FRAX 計算:https://frax.shef.ac.uk/FRAX/tool.aspx?lang=jp (国:日本)").font = Font(size=9, color="555555")
    ws.cell(r + 1, 2, "黄色セル=未入力/要入力。値は「はい/いいえ」をプルダウンで修正できます。").font = Font(size=9, color="555555")


def write_list(ws, results):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    headers = ["No", "受診日", "患者ID", "氏名", "年齢", "性別", "YAM(%)", "区分", "理由", "FRAXシート"]
    widths = [5, 12, 12, 16, 7, 7, 9, 9, 44, 26]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    n = {k: sum(1 for _, k2, _, _ in results if k2 == k) for k in ("対象", "要確認", "対象外")}
    ws["A1"] = f"FRAX 対象者判定一覧(対象 {n['対象']}名/要確認 {n['要確認']}名/対象外 {n['対象外']}名)"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = fill(NAVY)
    ws.row_dimensions[1].height = 26
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(2, c, h)
        cell.font = Font(bold=True)
        cell.fill = fill(BLUE)
        cell.border = BOX
        ws.column_dimensions[get_column_letter(c)].width = w
    # 対象 → 要確認 → 対象外 の順に並べる
    order = {"対象": 0, "要確認": 1, "対象外": 2}
    for i, (p, kind, reason, title) in enumerate(sorted(results, key=lambda x: order[x[1]]), 1):
        r = i + 2
        vals = [i, to_date(p.get("date")), p.get("pid"), p.get("name"), p.get("_age"), p.get("_sex"),
                to_num(p.get("yam")), kind, reason, title or ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.border = BOX
            cell.fill = fill({"対象": TARGET_FILL, "要確認": CHECK_FILL}.get(kind, EXCLUDE_FILL))
        ws.cell(r, 2).number_format = "yyyy/mm/dd"
        if title:
            link = ws.cell(r, 10)
            link.hyperlink = f"#'{title}'!A1"
            link.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{len(results) + 2}"


def build(src, out):
    rows, user_settings = read_rows(src)
    s = get_settings(user_settings)
    patients = load_patients(rows)
    wb = openpyxl.Workbook()
    lst = wb.active
    lst.title = "判定一覧"
    used, results = {"判定一覧"}, []
    for p in patients:
        kind, reason = judge(p, s)
        title = None
        if kind == "対象":
            title = sheet_title(p, used)
            write_frax_sheet(wb.create_sheet(title), p, s, reason)
        results.append((p, kind, reason, title))
    write_list(lst, results)
    wb.save(out)
    n = sum(1 for _, k, _, _ in results if k == "対象")
    print(f"{len(patients)}名中 {n}名が FRAX 対象 → {out}")
    for p, kind, reason, _ in results:
        print(f"  [{kind}] {p.get('pid')} {p.get('name')}:{reason}")


SAMPLES = [  # 架空の患者(動作確認用)
    ("2026-10-14", "S0001", "見本 花子", "女", "1961-04-02", 152, 48, 76, 0.62, "腰椎", "", "", "あり", "", "", "", "", ""),
    ("2026-10-14", "S0002", "見本 和子", "女", "1955-08-20", 148, 44, 65, 0.55, "腰椎", "", "", "", "", "", "", "", ""),
    ("2026-10-14", "S0003", "見本 太郎", "男", "1970-01-15", 170, 68, 92, 0.82, "腰椎", "", "", "", "あり", "", "", "", ""),
    ("2026-10-14", "S0004", "見本 良子", "女", "1949-11-30", 150, 52, 78, 0.60, "大腿骨", "", "あり", "", "", "", "", "", ""),
    ("2026-10-14", "S0005", "見本 春子", "女", "1952-03-03", 155, 50, 74, 0.61, "大腿骨", "", "", "", "", "あり", "", "", ""),
    ("2026-10-14", "S0006", "見本 次郎", "男", "1990-06-10", 175, 70, 72, 0.70, "腰椎", "", "", "", "", "", "", "", ""),
    ("2026-10-14", "S0007", "見本 秋子", "女", "1963-09-09", 158, 55, "", "", "", "", "", "", "", "", "", "", ""),
    ("2026-10-14", "S0008", "見本 冬子", "女", "1946-12-24", 145, 40, 79, 0.58, "大腿骨", "あり", "", "", "", "", "", "", ""),
]


def make_template(out):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "患者一覧"
    labels = [c[0] for c in COLUMNS if c[1] != "age"]  # 年齢は生年月日から計算(列を足せば年齢も可)
    for c, h in enumerate(labels, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill(NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = max(10, len(h) * 1.6)
    ws.row_dimensions[1].height = 36
    for r, row in enumerate(SAMPLES, 2):
        for c, v in enumerate(row, 1):
            if c in (1, 5) and v:
                v = dt.date.fromisoformat(v)
            cell = ws.cell(r, c, v if v != "" else None)
            if isinstance(v, dt.date):
                cell.number_format = "yyyy/mm/dd"
    dv = DataValidation(type="list", formula1='"あり,なし"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("K2:R1000")
    sx = DataValidation(type="list", formula1='"女,男"', allow_blank=True)
    ws.add_data_validation(sx)
    sx.add("D2:D1000")
    ws.freeze_panes = "D2"

    st = wb.create_sheet("設定")
    for c, (h, w) in enumerate(zip(["項目", "値", "説明"], [40, 12, 60]), 1):
        cell = st.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.fill = fill(BLUE)
        st.column_dimensions[get_column_letter(c)].width = w
    for r, (name, val, note) in enumerate(SETTINGS, 2):
        st.cell(r, 1, name)
        st.cell(r, 2, val).fill = fill(YELLOW)
        st.cell(r, 3, note).font = Font(size=9, color="555555")
    wb.save(out)
    print(f"テンプレートを作成 → {out}(サンプル患者は架空です。実データで上書きしてください)")


if __name__ == "__main__":
    args = sys.argv[1:]
    if len(args) == 2 and args[0] == "--template":
        make_template(args[1])
    elif len(args) in (1, 2):
        src = Path(args[0])
        build(src, args[1] if len(args) == 2 else src.with_name(f"FRAXシート_{src.stem}.xlsx"))
    else:
        sys.exit(__doc__)
